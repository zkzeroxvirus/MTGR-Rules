from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT
OUT = OUT_DIR / "MTGR-HOST-COMMAND-PROGRESSION-AND-SKILL-TREE-DESIGN-V0.2.docx"
TREE_IMAGE = OUT_DIR / "host-skill-tree-xml-mockup.png"

BLUE = "2E74B5"
NAVY = "1F355E"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF2CC"
PALE_GREEN = "E2F0D9"
PALE_RED = "FCE4D6"
PALE_PURPLE = "E4DFEC"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
BLACK = "111111"


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_inches, indent_dxa=120):
    widths_dxa = [round(v * 1440) for v in widths_inches]
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[idx])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row):
    repeat_header(row)


def set_cell_text(cell, text, bold=False, color=BLACK, size=8.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    set_font(r, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, value in enumerate(headers):
        shade_cell(hdr.cells[i], header_fill)
        set_cell_text(hdr.cells[i], value, bold=True, color=NAVY, size=font_size)
    repeat_header(hdr)
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        keep_row(row)
        if ridx % 2 == 1:
            for cell in row.cells:
                shade_cell(cell, "FAFAFA")
        for i, value in enumerate(row_data):
            set_cell_text(row.cells[i], value, size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label, text, fill=PALE_GOLD):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ": ")
    set_font(r, size=10, bold=True, color=NAVY)
    r = p.add_run(text)
    set_font(r, size=10, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_font(r, size=10.5)
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def configure_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 8),
        ("Subtitle", 14, GRAY, 0, 10),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in [s.name for s in doc.styles]:
        code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(7.5)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("MTGR Host Progression | Playtest Design V0.2")
    set_font(hr, size=8, color=GRAY, italic=True)
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("MTGR design proposal  |  ")
    set_font(fr, size=8, color=GRAY)
    add_page_number(fp)


def page_break(doc):
    doc.add_page_break()


def make_tree_mockup():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    w, h = 1800, 1050
    img = Image.new("RGB", (w, h), "#171A20")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 28)
        small = ImageFont.truetype("arial.ttf", 21)
        tiny = ImageFont.truetype("arial.ttf", 17)
        bold = ImageFont.truetype("arialbd.ttf", 29)
    except OSError:
        font = small = tiny = bold = ImageFont.load_default()

    d.rounded_rectangle((35, 30, 1765, 1015), radius=26, fill="#242833", outline="#8992A4", width=3)
    d.rounded_rectangle((65, 55, 1735, 140), radius=18, fill="#1F355E")
    d.text((95, 76), "HOST COMMANDS", font=bold, fill="white")
    d.text((1320, 79), "Talent 4 / 6     Rank 7", font=small, fill="#E8EEF5")

    tabs = [(75, "DOMINION", "#2E74B5"), (350, "TORMENT", "#8B2D52"), (625, "SUPPRESSION", "#3E6B4C")]
    for x, label, fill in tabs:
        d.rounded_rectangle((x, 160, x + 250, 225), radius=14, fill=fill)
        d.text((x + 38, 178), label, font=small, fill="white")
    d.rounded_rectangle((925, 160, 1715, 225), radius=14, fill="#343A46")
    d.text((955, 178), "Search / filter     Equipped: 2 of 3", font=small, fill="#DCE2EC")

    # left scrollable tree area
    d.rounded_rectangle((75, 250, 1235, 925), radius=18, fill="#1D212A", outline="#596273", width=2)
    d.text((105, 275), "DOMINION — Rule the Stack", font=bold, fill="#8FC7F2")
    d.text((105, 315), "Scrollable tree path • hover any node for full rules text", font=tiny, fill="#AAB3C2")
    nodes = [
        (180, 400, "TIER I", "Seal the Combo", "Rank 2 • 1 TP", "#2E74B5"),
        (690, 400, "TIER I", "Answer the Answer", "Rank 2 • 1 TP", "#2E74B5"),
        (180, 560, "TIER II", "Redirect the Assault", "Rank 4 • 1 TP", "#345F86"),
        (690, 560, "TIER II", "Close the Gate", "Rank 4 • 1 TP", "#345F86"),
        (435, 710, "TIER III", "Quell the Chorus", "Rank 7 • 1 TP", "#4B4D82"),
        (435, 835, "CAPSTONE", "Final Authority", "Rank 10 • 2 Authority", "#8B6B1F"),
    ]
    centers = []
    for x, y, tier, name, req, fill in nodes:
        box = (x, y, x + 365, y + 105)
        centers.append(((x + x + 365)//2, y, y+105))
        d.rounded_rectangle(box, radius=16, fill=fill, outline="#C5D3E6", width=2)
        d.text((x + 18, y + 12), tier, font=tiny, fill="#DDE8F7")
        d.text((x + 18, y + 39), name, font=small, fill="white")
        d.text((x + 18, y + 75), req, font=tiny, fill="#DDE8F7")
    # connectors
    lines = [((362,505),(362,560)),((872,505),(872,560)),((362,665),(617,710)),((872,665),(617,710)),((617,815),(617,835))]
    for a,b in lines:
        d.line((a,b), fill="#91A5BF", width=6)

    # right detail / equipment panel
    d.rounded_rectangle((1265, 250, 1715, 925), radius=18, fill="#202630", outline="#596273", width=2)
    d.text((1295, 280), "NODE DETAILS", font=bold, fill="white")
    d.text((1295, 330), "Final Authority", font=font, fill="#F2D47A")
    detail = [
        "Cost: 2 Authority",
        "Limit: once per encounter",
        "Unlock: Rank 10 + Quell the Chorus",
        "",
        "Counter target spell or ability.",
        "The spell or ability created by this",
        "Command can't be countered.",
        "",
        "Inspired by: Last Word, Disallow",
    ]
    y = 385
    for line in detail:
        d.text((1295, y), line, font=tiny, fill="#D7DCE5")
        y += 32
    d.rounded_rectangle((1295, 720, 1685, 785), radius=13, fill="#18652A")
    d.text((1390, 738), "EQUIP COMMAND", font=small, fill="white")
    d.rounded_rectangle((1295, 810, 1685, 885), radius=13, fill="#363D49")
    d.text((1320, 826), "Slot 1: Seal the Combo", font=tiny, fill="white")
    d.text((1320, 855), "Slot 2: Demonic Ward", font=tiny, fill="#C9D0DB")
    img.save(TREE_IMAGE)


LINEAR = [
    (1, "Demonic Ward", "Automatic at Rank 1", "1", "1/encounter", "Reaction", "Target Host permanent gains hexproof or indestructible until end of turn.", "Tamiyo's Safekeeping; Loran's Escape", "Generic defense"),
    (1, "Break the Charge", "Automatic at Rank 1", "1", "1/encounter", "After attackers", "Prevent all combat damage target attacking creature would deal this combat.", "Fog; Healing Salve", "Generic combat"),
    (1, "Rebuke", "Automatic at Rank 1", "1", "1/encounter", "Reaction", "Counter target spell unless its controller pays {2}.", "Quench; Mana Leak", "Dominion seed"),
    (1, "Painful Terms", "Automatic at Rank 1", "1", "1/encounter", "Reaction", "Target player may pay 2 life. If they don't, counter target activated or triggered ability they control.", "Dash Hopes; Stifle", "Torment seed"),
    (2, "Pin in Place", "Automatic at Rank 2", "1", "1/encounter", "Host priority", "Tap target nonland permanent controlled by a player. It doesn't untap during its controller's next untap step.", "Frost Breath; Frost Lynx", "Suppression seed"),
    (3, "Strategic Withdrawal", "Automatic at Rank 3", "1", "1/encounter", "Reaction", "Target Host permanent phases out until the beginning of the Host's next turn.", "Slip Out the Back; Teferi's Time Twist", "Generic preservation"),
    (5, "Reclaim the Fallen", "Automatic at Rank 5", "1", "1/encounter", "Host main phase", "Return target permanent card from the Host's graveyard to the Host's hand.", "Regrowth; Bala Ged Recovery", "Generic recovery"),
    (7, "Open the Way", "Automatic at Rank 7", "1", "1/encounter", "Before blockers", "Target Host creature gains trample until end of turn and can't be blocked by more than one creature this combat.", "Rancor; Bellowing Tanglewurm", "Generic offense"),
    (9, "Second Wind", "Automatic at Rank 9", "1", "1/encounter", "Host priority", "Untap up to two target nonland permanents controlled by the Host. Creatures untapped this way gain vigilance until end of turn.", "To Arms!; Vitalize", "Generic tempo"),
    (10, "Host's Decree", "Automatic at Rank 10", "1", "1/encounter", "Mode-dependent", "Choose one: grant a Host permanent hexproof and indestructible; counter a spell unless {4} is paid; tap up to two player-controlled nonland permanents; or return a Host permanent card from the graveyard to hand.", "Cryptic Command; Command cycle", "Linear capstone"),
]


BRANCHES = {
    "Dominion": [
        ("Tier I-A", "Seal the Combo", "Rank 2; 1 TP", "1", "1/encounter", "Reaction", "Counter target activated or triggered ability. If its source is a permanent, its nonmana activated abilities can't be activated until end of turn.", "Stifle; Disallow; Pithing Needle"),
        ("Tier I-B", "Answer the Answer", "Rank 2; 1 TP", "1", "1/encounter", "Reaction", "Counter target spell that targets a spell or ability controlled by the Host.", "Turn Aside; Dovin's Veto"),
        ("Tier II-A", "Redirect the Assault", "Rank 4; Seal the Combo; 1 TP", "1", "1/encounter", "Reaction", "Change the target of target spell or ability with a single target. A different legal target must exist.", "Misdirection; Deflecting Swat"),
        ("Tier II-B", "Close the Gate", "Rank 4; Answer the Answer; 1 TP", "1", "1/encounter", "Reaction", "Counter target spell cast from anywhere other than its controller's hand.", "Drannith Magistrate; Grafdigger's Cage"),
        ("Tier III", "Quell the Chorus", "Rank 7; any Tier II Dominion node; 1 TP", "1", "1/encounter", "Reaction", "Counter any number of target triggered abilities controlled by players that share a source or trigger event.", "Summary Dismissal; Whirlwind Denial"),
        ("Capstone", "Final Authority", "Rank 10; Quell the Chorus; 1 TP", "2", "1/encounter", "Reaction", "Counter target spell or ability. The spell or ability created by this Command can't be countered.", "Last Word; Disallow; Dovin's Veto"),
    ],
    "Torment": [
        ("Tier I-A", "Cruel Choice", "Rank 2; 1 TP", "1", "1/encounter", "Second-spell trigger", "When a player casts their second spell in a turn, they may pay 5 life. If they don't, counter that spell.", "Dash Hopes; Archon of Emeria"),
        ("Tier I-B", "Forced Exposure", "Rank 2; 1 TP", "1", "1/encounter", "Host main phase", "Target player reveals their hand. Choose a nonland card. They may discard it; if they don't, spells with that name cost {3} more until the end of their next turn.", "Thoughtseize; Elite Spellbinder; Nevermore"),
        ("Tier II-A", "Price of Ambition", "Rank 4; Cruel Choice; 1 TP", "1", "1/encounter", "Host priority", "Until end of turn, the next time target player casts a spell, they discard a card or lose 5 life.", "Painful Quandary"),
        ("Tier II-B", "Divide Their Loyalty", "Rank 4; Forced Exposure; 1 TP", "1", "1/encounter", "Host main phase", "Separate target player's creatures into two piles. That player chooses one pile and sacrifices those creatures.", "Do or Die; Make an Example"),
        ("Tier III", "No Safe Answer", "Rank 7; any Tier II Torment node; 1 TP", "2", "1/encounter", "Host main phase", "Target player repeats twice: discard a card, sacrifice a nonland permanent, or lose 3 life.", "Torment of Hailfire"),
        ("Capstone", "Marked for Death", "Rank 10; No Safe Answer; 1 TP", "2", "1/encounter", "Host priority", "Mark target player until the Host's next turn. The first time they decline to pay a life, mana, discard, or sacrifice option from a Host Command, double that Command's stated default life loss or increase its nonlife default by one object.", "Captive Audience; curse effects"),
    ],
    "Suppression": [
        ("Tier I-A", "Strip the Shield", "Rank 2; 1 TP", "1", "1/encounter", "Host priority", "Permanents controlled by target player lose hexproof and indestructible until end of turn.", "Shadowspear"),
        ("Tier I-B", "Grave Interdict", "Rank 2; 1 TP", "1", "1/encounter", "Reaction or priority", "Exile target card from a graveyard. Until end of turn, players can't cast spells from graveyards or activate abilities of cards in graveyards.", "Soul-Guide Lantern; Ground Seal"),
        ("Tier II-A", "Silence the Engine", "Rank 4; Strip the Shield; 1 TP", "1", "1/encounter", "Host priority", "Nonmana activated abilities of target nonland permanent can't be activated until the Host's next turn.", "Linvala, Keeper of Silence; Cursed Totem"),
        ("Tier II-B", "Seal the Threshold", "Rank 4; Grave Interdict; 1 TP", "1", "1/encounter", "Host priority", "Until end of turn, creatures entering or dying don't cause abilities controlled by players to trigger.", "Hushbringer; Torpor Orb"),
        ("Tier III", "Divide and Conquer", "Rank 7; any Tier II Suppression node; 1 TP", "2", "1/encounter", "Host main phase", "Tap all creatures target player controls. Choose up to two of them; those creatures don't untap during that player's next untap step.", "Sleep; Kiora Bests the Sea God"),
        ("Capstone", "Absolute Suppression", "Rank 10; Divide and Conquer; 1 TP", "2", "1/encounter", "Host priority", "Until end of turn, permanents controlled by players lose hexproof and indestructible, and their nonmana activated abilities can't be activated.", "Shadowspear; Dress Down; Linvala"),
    ],
}


ACHIEVEMENTS = [
    ("Turn the Blade", "Dominion", "Redirect a player-controlled spell or ability into another player-controlled object in three verified encounters.", "1", "1/encounter", "Change the target of a single-target spell or ability; if it now targets its controller's object, copy it and the Host chooses the copy's target.", "Deflecting Swat; Reiterate"),
    ("Chainbreaker", "Dominion", "Counter three or more objects on the stack across one verified run.", "2", "1/encounter", "Exile target spell and each copy of it on the stack.", "Mindbreak Trap"),
    ("No Further Appeals", "Dominion", "Win a verified encounter after a player casts four or more spells in one turn.", "1", "1/encounter", "After the current spell resolves, players can't cast spells for the rest of the turn.", "Silence; Orim's Chant"),
    ("Unmake the Invocation", "Dominion", "Counter a spell cast from a graveyard, exile, or without mana in two verified encounters.", "2", "1/encounter", "Copy target instant or sorcery controlled by a player, then return it to its owner's hand. The Host chooses new targets for the copy.", "Narset's Reversal"),
    ("Devil's Bargain", "Torment", "In one verified run, have players choose both sides of Torment Commands at least once.", "1", "1/encounter", "Target player chooses: the Host creates a 4/4 red Devil token with haste, or that player loses 4 life.", "Vexing Devil"),
    ("Burn the Evidence", "Torment", "Cause players to discard or mill at least fifteen total cards in a verified run without eliminating a player.", "1", "1/encounter", "Target player chooses: mill ten cards, or Burn the Evidence deals 5 damage to them.", "Book Burning"),
    ("Kneel or Bleed", "Torment", "Have three player creatures die in one turn after their controller declined a Command payment.", "2", "1/encounter", "For each creature target player controls, that player pays 2 life or sacrifices that creature.", "Killing Wave"),
    ("Ransom the Fallen", "Torment", "Resolve three different Torment Commands during one verified run.", "1", "1/encounter", "Target player sacrifices a creature unless they discard two cards.", "Soul Shatter; Tribute to Hunger"),
    ("Faltering Advance", "Suppression", "Prevent at least ten combat damage using Commands across a verified run.", "1", "1/encounter", "Tap up to two target creatures. They don't untap during their controller's next untap step.", "Frost Breath"),
    ("Strip Their Purpose", "Suppression", "Disable three different kinds of permanent in one verified run.", "1", "1/encounter", "Target creature loses all abilities and has base power and toughness 1/1 until end of turn.", "Humble; Dress Down"),
    ("Contain the Intruder", "Suppression", "Exile or stop three cards attempting to enter from graveyards or libraries in one verified run.", "1", "1/encounter", "Until end of turn, if a nontoken creature would enter under a player's control and it wasn't cast, exile it instead.", "Containment Priest"),
    ("Narrow the Battlefield", "Suppression", "Survive a verified encounter in which players controlled twelve or more creatures at once.", "2", "1/encounter", "Until the Host's next turn, each player can't attack with more than two creatures each combat.", "Crawlspace; Silent Arbiter"),
]


RANKS = [
    (1, 0, "Demonic Ward; Break the Charge; Rebuke; Painful Terms", "Default Loadout 1"),
    (2, 100, "Pin in Place; +1 Talent Point", "Tier I access"),
    (3, 250, "Strategic Withdrawal", "Saved Loadout Preset 2"),
    (4, 450, "+1 Talent Point", "Tier II access"),
    (5, 700, "Reclaim the Fallen", "Host title milestone"),
    (6, 1000, "+1 Talent Point", "Build breadth milestone"),
    (7, 1350, "Open the Way; +1 Talent Point", "Tier III access; Preset 3"),
    (8, 1750, "+1 Talent Point", "Profile treatment milestone"),
    (9, 2200, "Second Wind", "Veteran Host title"),
    (10, 2700, "Host's Decree; +1 Talent Point", "Capstone access; capstone cosmetic"),
]


XML_SAMPLE = dedent("""
<Panel id="hostProgression" active="false" width="1180" height="760"
       allowDragging="true" raycastTarget="true">
  <HorizontalLayout spacing="8" padding="12 12 12 12">
    <Button id="tabDominion" onClick="showBranch(Dominion)">DOMINION</Button>
    <Button id="tabTorment" onClick="showBranch(Torment)">TORMENT</Button>
    <Button id="tabSuppression" onClick="showBranch(Suppression)">SUPPRESSION</Button>
    <Button id="tabLinear" onClick="showBranch(Linear)">RANK REWARDS</Button>
  </HorizontalLayout>

  <HorizontalLayout spacing="12" padding="12 12 72 12">
    <ScrollView id="treeScroll" width="760" height="590"
                verticalScrollbarVisibility="Auto" horizontalScrollbarVisibility="Auto">
      <Panel id="branchCanvas" width="1100" height="900">
        <!-- Lua creates NodeButton panels from managed Command definitions. -->
        <Button id="node_dominion_1a" rectAlignment="UpperLeft"
                offsetXY="90 -100" width="260" height="82"
                onClick="selectNode(dominion_1a)" tooltip="Seal the Combo" />
        <Button id="node_dominion_2a" rectAlignment="UpperLeft"
                offsetXY="90 -300" width="260" height="82"
                onClick="selectNode(dominion_2a)" tooltip="Redirect the Assault" />
      </Panel>
    </ScrollView>

    <Panel id="nodeDetails" width="360" height="590" raycastTarget="true">
      <Text id="detailName" fontSize="26" />
      <Text id="detailUnlock" fontSize="16" />
      <Text id="detailCost" fontSize="16" />
      <Text id="detailRules" fontSize="17" />
      <Button id="equipSelected" onClick="equipSelectedCommand">EQUIP</Button>
    </Panel>
  </HorizontalLayout>

  <Panel id="equipmentTray" rectAlignment="LowerCenter" height="64"
         raycastTarget="true">
    <Text id="slot1">Slot 1</Text><Text id="slot2">Slot 2</Text>
    <Text id="slot3">Slot 3</Text><Button onClick="closeHostProgression">CLOSE</Button>
  </Panel>
</Panel>
""").strip()


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(92)
    r = p.add_run("MTGR")
    set_font(r, size=13, bold=True, color=BLUE)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Host Command Progression\nand Skill Tree Design")
    set_font(r, size=28, bold=True, color=NAVY)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A complete playtest framework for linear unlocks, branch specialization, achievements, Command economy, and Tabletop Simulator XML")
    set_font(r, size=14, color=GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    r = p.add_run("DESIGN PROPOSAL V0.2  |  AUGUST 2026")
    set_font(r, size=10, bold=True, color="8B6B1F")
    add_callout(doc, "Status", "Everything in this document is a playtest proposal. Names, costs, thresholds, effects, and achievement requirements remain subject to dev-team review and live testing.", PALE_GOLD)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Design objective")
    set_font(r, size=11, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Every Host receives a useful core toolkit. Branch investment changes how the Host exerts pressure; it does not merely increase raw statistics.")
    set_font(r, size=12, italic=True, color=GRAY)


def add_contents(doc):
    page_break(doc)
    doc.add_heading("Contents and decision summary", level=1)
    rows = [
        ("1", "Progression architecture", "Linear ranks, talent points, achievements, capstones"),
        ("2", "Command economy", "Authority cost, equipped slots, encounter invocation limits"),
        ("3", "Rank 1-10 advancement", "D&D-style progression table"),
        ("4", "Linear Command register", "Ten automatic unlocks"),
        ("5", "Skill trees", "Dominion, Torment, Suppression; six nodes each"),
        ("6", "Achievement Commands", "Twelve optional sidegrades"),
        ("7", "TTS XML representation", "Tabbed and scrollable interface concept"),
        ("8", "Implementation and playtest", "Data contract, telemetry, unresolved decisions"),
    ]
    add_table(doc, ["Section", "Topic", "Purpose"], rows, [0.65, 2.2, 3.65], font_size=9.5)
    add_callout(doc, "Proposed roster", "40 total Commands: 10 linear, 18 skill-tree, and 12 achievement Commands. The three branch capstones are included among the 18 skill-tree Commands.", PALE_BLUE)
    add_bullet(doc, "Linear progression supplies generic utility and low-power introductions to Dominion, Torment, and Suppression.")
    add_bullet(doc, "Talent Points unlock stronger branch-specific sidegrades; six points are available by Rank 10.")
    add_bullet(doc, "Achievements unlock specialized Commands but never additional Authority, slots, or encounter uses.")
    add_bullet(doc, "Every Command cites recognizable Magic card inspirations and uses familiar Magic templating where practical.")


def add_architecture(doc):
    page_break(doc)
    doc.add_heading("1. Progression architecture", level=1)
    add_callout(doc, "Core principle", "Rank grants breadth. Skill-tree investment grants identity. Achievements grant unusual sidegrades. Capstones grant a branch-defining moment, not a permanent increase to Host statistics.", PALE_GREEN)
    doc.add_heading("Four unlock channels", level=2)
    rows = [
        ("Linear Rank", "Automatic at listed Host Rank", "Generic utility plus weaker branch foundations", "10"),
        ("Skill Tree", "Spend 1 Talent Point and meet rank/prerequisites", "Stronger specialization and loadout identity", "18"),
        ("Achievement", "Complete the listed verified Host achievement", "Unusual sidegrades and mastery rewards", "12"),
        ("Capstone", "Rank 10, prerequisite mastery node, and 1 Talent Point", "Maximum branch expression; costs 2 Authority", "3 included above"),
    ]
    add_table(doc, ["Channel", "Unlock method", "Design purpose", "Count"], rows, [1.1, 1.75, 2.85, 0.8], font_size=8.8)
    doc.add_heading("Talent Point rules", level=2)
    for text in [
        "Talent Points are earned at Ranks 2, 4, 6, 7, 8, and 10: six points total.",
        "Each node costs 1 Talent Point. A Host may invest across multiple branches.",
        "Tier I requires Rank 2. Tier II requires Rank 4 and its linked Tier I node. Tier III requires Rank 7 and any Tier II node in that branch.",
        "A capstone requires Rank 10, its Tier III mastery node, and 1 Talent Point.",
        "During the playtest, full respecs remain free outside an active or pending-verification run.",
        "Unlocking a Command adds it to the library; it does not equip it or add a Command Slot.",
    ]:
        add_bullet(doc, text)
    doc.add_heading("Command resolution modes", level=2)
    rows = [
        ("command_effect", "The invocation is a Host system action and is not a Magic spell or ability. Generic counterspells cannot target it; players may still alter the game state or invalidate targets."),
        ("casts_spell", "The Command creates or copies a Magic spell. That spell uses the stack, may be responded to, and may be countered unless stated otherwise."),
    ]
    add_table(doc, ["Mode", "Rules behavior"], rows, [1.35, 5.15], font_size=9.2)


def add_economy(doc):
    page_break(doc)
    doc.add_heading("2. Command economy and encounter limits", level=1)
    add_callout(doc, "Recommendation", "Guarantee at least two potential Command interactions in every encounter. Keep Authority as a shared run resource, keep each equipped Command once per encounter, and allow a third invocation only in encounters with four or more non-Host players.", PALE_GOLD)
    doc.add_heading("Proposed scaling", level=2)
    rows = [
        (1, 2, 3, 2), (2, 2, 4, 2), (3, 3, 5, 2),
        (4, 3, 6, 3), (5, 4, 7, 3), (6, 4, 8, 3),
    ]
    add_table(doc, ["Non-Host players", "Equipped slots", "Starting Authority / run", "Invocation limit / encounter"], rows, [1.35, 1.35, 1.9, 1.9], font_size=9.2)
    doc.add_heading("Cost and limit standards", level=2)
    rows = [
        ("Standard Command", "1 Authority", "Once per encounter", "Most linear, Tier I-II, and achievement sidegrades"),
        ("Major Command", "2 Authority", "Once per encounter", "Wide battlefield effects, repeated Torment choices, stack exile"),
        ("Capstone", "2 Authority", "Once per encounter", "Branch-defining maximum effect"),
        ("Invocation limit", "No Authority cost", "2 or 3 total per encounter", "Prevents many equipped options from becoming excessive action density"),
    ]
    add_table(doc, ["Category", "Cost", "Default use limit", "Purpose"], rows, [1.25, 1.05, 1.45, 2.75], font_size=9)
    page_break(doc)
    doc.add_heading("Legacy power comparison", level=2)
    rows = [
        ("Not Today", "One encounter-refreshing sacrifice effect that counters any spell, activated ability, or triggered ability.", "Equivalent to a broad, zero-Authority interaction."),
        ("Absolute Command", "Once per encounter, cast a free copy of Disallow.", "A second broad interaction that uses the stack."),
        ("Proposed baseline", "Two equipped uses per encounter, usually narrower and costing 1 shared Authority each.", "At least the same interaction count, but with visible loadout choices and run-level resource pressure."),
        ("Large encounters", "Up to three invocations with four or more non-Host players.", "Adds scaling without granting a permanent Rank-based frequency advantage."),
    ]
    add_table(doc, ["System", "Strength", "Design implication"], rows, [1.25, 2.75, 2.5], font_size=8.8)
    add_callout(doc, "Non-stacking rule", "When Commands are enabled, do not also create Not Today tokens or grant legacy free Disallow uses. Passive Host Scaling remains active.", PALE_RED)


def add_rank_table(doc):
    page_break(doc)
    doc.add_heading("3. Host advancement by Rank", level=1)
    p = doc.add_paragraph("This table is the primary player-facing progression reference, modeled after a tabletop RPG class-level table.")
    p.paragraph_format.space_after = Pt(10)
    rows = [(r, f"{h:,}", commands, features) for r, h, commands, features in RANKS]
    add_table(doc, ["Rank", "Minimum HXP", "Command / Talent reward", "Other features"], rows, [0.65, 1.05, 3.1, 1.7], font_size=8.7)
    add_callout(doc, "Frequency protection", "Rank does not increase Starting Authority, equipped slots, or the encounter invocation limit. Those values scale only with the number of non-Host players.", PALE_BLUE)


def add_linear(doc):
    page_break(doc)
    doc.add_heading("4. Linear Command register", level=1)
    doc.add_paragraph("Linear Commands are automatically available at the listed Rank. They provide a viable general-purpose toolkit without matching the precision or ceiling of specialized branch Commands.")
    rows = []
    for rank, name, unlock, cost, limit, timing, effect, inspiration, role in LINEAR:
        rows.append((f"R{rank}\n{name}", f"{cost} A\n{limit}", timing, effect, inspiration, role))
    add_table(doc, ["Unlock / Command", "Cost / limit", "Timing", "Effect", "Magic inspiration", "Role"], rows, [1.0, 0.7, 0.8, 2.25, 1.05, 0.7], font_size=7.5)
    add_callout(doc, "Host's Decree", "Its Rank 10 reward is flexibility, not raw superiority. Every individual mode remains below the corresponding branch capstone.", PALE_GOLD)


def add_branch(doc, name, color_fill, philosophy, identity):
    page_break(doc)
    doc.add_heading(f"5.{ {'Dominion':'1','Torment':'2','Suppression':'3'}[name] } {name} skill tree", level=1)
    add_callout(doc, "Philosophy", philosophy, color_fill)
    doc.add_paragraph(identity)
    visual_rows = [
        ("Tier I", BRANCHES[name][0][1], BRANCHES[name][1][1], "Rank 2"),
        ("Tier II", BRANCHES[name][2][1], BRANCHES[name][3][1], "Rank 4 + linked Tier I"),
        ("Tier III", BRANCHES[name][4][1], "", "Rank 7 + any Tier II"),
        ("Capstone", BRANCHES[name][5][1], "", "Rank 10 + Tier III"),
    ]
    add_table(doc, ["Path", "Left specialization", "Right specialization", "Gate"], visual_rows, [1.05, 2.0, 2.0, 1.45], header_fill=color_fill, font_size=9)
    doc.add_heading(f"{name} Command register", level=2)
    rows = []
    for node, cmd, unlock, cost, limit, timing, effect, inspiration in BRANCHES[name]:
        rows.append((f"{node}\n{cmd}", unlock, f"{cost} A\n{limit}", timing, effect, inspiration))
    add_table(doc, ["Node / Command", "Unlock", "Cost / limit", "Timing", "Effect", "Magic inspiration"], rows, [1.05, 1.2, 0.72, 0.73, 1.8, 1.0], header_fill=color_fill, font_size=7.6)


def add_achievements(doc):
    page_break(doc)
    doc.add_heading("6. Achievement Command register", level=1)
    doc.add_paragraph("Achievement Commands are permanent library unlocks. They are optional sidegrades: achievements never grant additional Authority, Command Slots, or encounter invocations.")
    rows = []
    for name, branch, unlock, cost, limit, effect, inspiration in ACHIEVEMENTS:
        rows.append((name, branch, unlock, f"{cost} A\n{limit}", effect, inspiration))
    add_table(doc, ["Command", "Affinity", "Verified unlock condition", "Cost / limit", "Effect", "Magic inspiration"], rows, [0.95, 0.65, 1.6, 0.68, 1.72, 0.9], font_size=7.3)
    add_callout(doc, "Verification", "Achievement progress is server-owned and credited only from verified runs or an auditable moderator override. A Command unlock must be idempotent.", PALE_BLUE)


def add_xml(doc):
    page_break(doc)
    doc.add_heading("7. Tabletop Simulator XML representation", level=1)
    doc.add_paragraph("The Host Profile object should open a local, player-facing UI that visually matches the Player Profile. The tree is data-driven: managed Lua supplies node state and rules text, while XML supplies stable panels, scroll regions, tabs, detail views, and equipment slots.")
    doc.add_picture(str(TREE_IMAGE), width=Inches(6.5))
    cap = doc.add_paragraph("Figure 1. Proposed object-local XML skill-tree view. The active branch is tabbed, the tree canvas scrolls, the detail panel intercepts clicks, and equipment remains visible.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        set_font(run, size=8.5, italic=True, color=GRAY)
    doc.add_heading("Interaction requirements", level=2)
    for text in [
        "Tabs: Dominion, Torment, Suppression, Rank Rewards, Achievements, and Equipped.",
        "Scrollable branch canvas: show the complete path, including locked nodes and prerequisite connectors.",
        "Tooltips: concise summary on hover; full rules text, inspiration, cost, limit, and unlock condition in the detail panel.",
        "State treatment: locked, available, learned, equipped, and unavailable-during-run must be visually distinct.",
        "Click safety: every overlay panel and button uses raycastTarget=true so clicks do not pass through to object buttons.",
        "Orientation: copy the Player Profile object's proven transform/orientation strategy rather than compensating with mirrored text or reversed XML layout.",
        "Command equipment behaves like the Player Profile Essence page: browse in a dedicated local panel, select a Command, then assign it to a visible slot.",
    ]:
        add_bullet(doc, text)
    page_break(doc)
    doc.add_heading("Representative XML skeleton", level=2)
    for line in XML_SAMPLE.splitlines():
        p = doc.add_paragraph(style="Code Block")
        r = p.add_run(line if line else " ")
        set_font(r, name="Consolas", size=7.5, color=BLACK)
        if line.strip().startswith("<Panel") or line.strip().startswith("</Panel"):
            p.paragraph_format.keep_with_next = False


def add_implementation(doc):
    page_break(doc)
    doc.add_heading("8. Implementation and playtest contract", level=1)
    doc.add_heading("Managed Command definition", level=2)
    rows = [
        ("id", "Stable machine identifier; never use the display name as the save key."),
        ("name", "Player-facing Command name."),
        ("source", "linear, skill_tree, achievement, or legacy migration."),
        ("branch", "dominion, torment, suppression, or neutral."),
        ("tier / rank", "Tree tier or automatic Rank gate."),
        ("prerequisites", "Node IDs, Rank, achievement ID, and required Talent Points."),
        ("authorityCost", "Normally 1; major Commands and capstones cost 2."),
        ("perEncounterLimit", "Normally 1."),
        ("resolutionMode", "command_effect or casts_spell."),
        ("timing", "Exact legal invocation window."),
        ("effectText", "Complete player-facing rules text."),
        ("inspiration", "Magic cards used as the mechanical foundation."),
        ("tracking", "Marker/token requirements and expiration cleanup."),
        ("version", "Managed definition version for migration and audit."),
    ]
    add_table(doc, ["Field", "Requirement"], rows, [1.4, 5.1], font_size=8.8)
    doc.add_heading("Run lock and equipment rules", level=2)
    for text in [
        "Loadouts lock when a run becomes active or pending verification.",
        "An equipped Command must be unlocked, valid for the current definition version, and occupy exactly one slot.",
        "If a respec removes an equipped Command, the loadout becomes invalid and must be repaired before the next run locks.",
        "Equipped Commands, remaining Authority, per-Command use state, and remaining encounter invocations are public information.",
        "Achievement and Rank unlocks are account-wide Host Profile data; encounter use state is session data.",
    ]:
        add_bullet(doc, text)
    doc.add_heading("Playtest metrics", level=2)
    rows = [
        ("Interaction floor", "Host can legally make at least two Command interactions in a typical encounter when Authority remains."),
        ("Authority pressure", "Hosts cannot use their preferred maximum number in every encounter without considering the remaining run."),
        ("Branch identity", "Testers can identify the branch from the effect without seeing its label."),
        ("Counterplay", "Players report a meaningful sequencing, payment, protection, or timing response."),
        ("Resolution time", "A Command normally resolves faster than the player action it answers."),
        ("Selection diversity", "No Command dominates equipped slots across most decks and player counts."),
        ("Tracking burden", "Temporary effects expire cleanly and require no private Host notes."),
    ]
    add_table(doc, ["Measure", "Pass condition"], rows, [1.45, 5.05], font_size=9)
    add_callout(doc, "First balance checkpoint", "Playtest the 10 linear Commands plus Tier I nodes before implementing every advanced and achievement Command. The complete roster defines the destination; it does not require a single all-at-once release.", PALE_GREEN)


def add_sources(doc):
    page_break(doc)
    doc.add_heading("Appendix A. Design sources and precedents", level=1)
    doc.add_paragraph("Primary MTGR rules references used for this proposal:")
    for text in [
        "rules/host/authority-level.md — Starting Authority, Command Slots, run resource behavior, and legacy non-stacking rule.",
        "rules/host/authority-interactions.md — command_effect and casts_spell resolution modes.",
        "rules/host/host-profile-progression.md — Rank thresholds and Talent Point awards.",
        "rules/host/host-skill-tree.md — branch philosophies, tier gates, and design guardrails.",
        "rules/host/host-commands.md — current starter and skill-tree Commands.",
        "archive/legacy-sources/HOST-AUTHORITY-SYSTEM-V1.0.md — Not Today and Absolute Command power benchmark.",
    ]:
        add_bullet(doc, text)
    doc.add_paragraph("Magic design precedents:")
    for text in [
        "Wizards of the Coast, Mechanical Color Pie 2021 and Mechanical Color Pie 2021 Changes.",
        "Wizards of the Coast, A La Mode — opponent-choice and punisher design.",
        "Individual Magic cards listed in each Command register as mechanical inspirations; the proposed Commands are adaptations, not reproductions of card identity or artwork.",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "Terminology note", "The prior branch name Control is replaced by Suppression throughout this proposal because control is a broad Magic strategy encompassing all three branches. Suppression more precisely describes the branch's battlefield and zone-denial effects.", PALE_PURPLE)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    make_tree_mockup()
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "MTGR Host Command Progression and Skill Tree Design V0.2"
    doc.core_properties.subject = "Host progression, Commands, skill trees, achievements, and TTS XML"
    doc.core_properties.author = "MTGR Development Team"
    doc.core_properties.keywords = "MTGR, Host, Commands, progression, skill tree, Tabletop Simulator"
    add_cover(doc)
    add_contents(doc)
    add_architecture(doc)
    add_economy(doc)
    add_rank_table(doc)
    add_linear(doc)
    add_branch(doc, "Dominion", PALE_BLUE, "Rule the stack. Dominion decides whether, how, and against what spells and abilities resolve.", "Identity test: if the Command's principal job is to determine the outcome or targeting of an object on the stack, it belongs in Dominion.")
    add_branch(doc, "Torment", PALE_RED, "Make every choice hurt. Torment converts ambition and incomplete information into life, card, mana, or permanent decisions.", "Identity test: if the Command's principal job is to make a player choose how they pay for continuing their plan, it belongs in Torment.")
    add_branch(doc, "Suppression", PALE_GREEN, "Reshape the battlefield. Suppression disables protections, engines, combatants, and zones after they exist.", "Identity test: if the Command's principal job is to change what permanents, zones, or combatants can do, it belongs in Suppression.")
    add_achievements(doc)
    add_xml(doc)
    add_implementation(doc)
    add_sources(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
