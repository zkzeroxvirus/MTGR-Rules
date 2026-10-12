from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


OUT = Path(__file__).resolve().parents[1] / "MTGR-HOST-COMMAND-BRANCH-DESIGN-GUIDE-V0.1.docx"
BLUE = "2E74B5"
NAVY = "1F355E"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF2CC"
PALE_GREEN = "E2F0D9"
PALE_RED = "FCE4D6"
GRAY = "666666"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_fixed(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_status_box(doc, label, text, fill=PALE_GOLD):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_width(cell, 6.5)
    set_cell_margins(cell, 120, 160, 120, 160)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_label_para(doc, label, text, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_command(doc, name, branch, status, timing, cost, rules, identity, counterplay, tracking, abuse):
    h = doc.add_heading(name, level=3)
    keep_with_next(h)
    status_fill = PALE_GREEN if status == "APPROVED PLAYTEST" else PALE_GOLD if status == "DESIGN CONCEPT" else PALE_RED
    table = doc.add_table(rows=1, cols=2)
    set_table_fixed(table, [1.65, 4.85])
    shade_cell(table.cell(0, 0), status_fill)
    shade_cell(table.cell(0, 1), PALE_BLUE)
    table.cell(0, 0).paragraphs[0].add_run(status).bold = True
    table.cell(0, 1).paragraphs[0].add_run(branch).bold = True
    for cell in table.rows[0].cells:
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    fields = (
        ("Timing", timing),
        ("Cost / limit", cost),
        ("Rules text", rules),
        ("Branch fit", identity),
        ("Counterplay", counterplay),
        ("Tracking", tracking),
        ("Abuse check", abuse),
    )
    for index, (label, text) in enumerate(fields):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if index < len(fields) - 1:
            p.paragraph_format.keep_with_next = True
        r = p.add_run(label + ": ")
        r.bold = True
        p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, before, after in (
    ("Title", 26, NAVY, 0, 10),
    ("Subtitle", 13, GRAY, 0, 16),
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, "1F4D78", 10, 5),
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    if name.startswith("Heading"):
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.left_indent = Inches(0)
        style.paragraph_format.first_line_indent = Inches(0)

for list_name in ("List Bullet", "List Number"):
    style = styles[list_name]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.375)
    style.paragraph_format.first_line_indent = Inches(-0.188)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

header = section.header
hp = header.paragraphs[0]
hp.text = "MTGR HOST COMMAND DESIGN  |  DEVELOPMENT DRAFT"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.runs[0].font.size = Pt(8)
hp.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("MTGR Host Command Branch Design Guide  |  V0.1  |  Page ")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string(GRAY)
add_page_field(fp)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(42)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("MTGR")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(BLUE)

title = doc.add_paragraph(style="Title")
title.add_run("Host Command Branch\nDesign Guide")
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.add_run("Branch philosophy, mechanical boundaries, and example Commands")

line = doc.add_paragraph()
line.paragraph_format.space_after = Pt(14)
run = line.add_run("DEVELOPMENT DRAFT  /  VERSION 0.1  /  AUGUST 30, 2026")
run.bold = True
run.font.size = Pt(9)
run.font.color.rgb = RGBColor.from_string(GRAY)

add_status_box(
    doc,
    "Purpose",
    "Give the design and development teams a shared jumping-off point for building Host Commands without blurring branch identity or overwhelming players.",
    PALE_BLUE,
)

doc.add_heading("Design north star", level=1)
p = doc.add_paragraph()
r = p.add_run("The Host should feel powerful because they make decisive, legible interventions - not because they invalidate the players' game.")
r.bold = True
p = doc.add_paragraph(
    "Every Command should create a short tactical problem with a visible trigger, a clear answer window, and minimal bookkeeping. "
    "Commands are sidegrades that change how the Host exercises Authority; they are not passive statistical upgrades."
)

doc.add_heading("Status language", level=2)
table = doc.add_table(rows=4, cols=2)
table.style = "Table Grid"
set_table_fixed(table, [1.65, 4.85])
for cell in table.rows[0].cells:
    shade_cell(cell, PALE_BLUE)
table.cell(0, 0).text = "Status"
table.cell(0, 1).text = "Meaning"
for cell in table.rows[0].cells:
    cell.paragraphs[0].runs[0].bold = True
rows = [
    ("APPROVED PLAYTEST", "Already defined for current playtesting. Wording may still be tuned."),
    ("DESIGN CONCEPT", "A starting point for discussion. Not approved and not yet wired into the game."),
    ("PLACEHOLDER / STUB", "A reserved node or name without approved behavior."),
]
for idx, (a, b) in enumerate(rows, 1):
    table.cell(idx, 0).text = a
    table.cell(idx, 1).text = b
set_repeat_table_header(table.rows[0])

doc.add_page_break()

# System overview
doc.add_heading("1. System Snapshot", level=1)
add_label_para(doc, "Host progression", "Linear Rank progression supplies the baseline toolkit. Skill-tree branches and achievements supply specialization and expression.")
add_label_para(doc, "Command loadout", "The Host equips a limited selection of unlocked Commands from the physical Host Profile object.")
add_label_para(doc, "Table representation", "Each equipped Command can spawn an Infinite Bag containing its Command card/token. Final art and object asset IDs will be supplied later.")
add_label_para(doc, "Target scale", "Approximately 30 total powers: 6 Rank, 18 branch, and 6 achievement powers.")

doc.add_heading("What a Command is", level=2)
add_bullets(doc, [
    "A discrete Host action with a known timing window, legal targets, cost, and limit.",
    "A tactical sidegrade that changes decision-making rather than increasing raw resources.",
    "A table-readable event that should resolve quickly and leave little persistent bookkeeping.",
])

doc.add_heading("Shared guardrails", level=2)
add_bullets(doc, [
    "Do not grant raw life, damage, card draw, lands, Authority, talent points, or Command slots.",
    "The trigger and legal targets must be known before the Host spends the Command.",
    "The affected player must have meaningful counterplay: sequencing, payment, target choice, stack response, or avoidance.",
    "Avoid private notes, hidden durations, and effects that require the table to remember several state changes.",
    "A Command should not perform another branch's defining job more efficiently.",
    "The strongest use should be exciting once per encounter and unacceptable if repeated freely.",
])

doc.add_heading("Branch map", level=2)
table = doc.add_table(rows=4, cols=5)
table.style = "Table Grid"
set_table_fixed(table, [1.05, 1.25, 1.55, 1.35, 1.30])
headers = ["Branch", "Promise", "Primary target", "Core verbs", "Avoid"]
for i, text in enumerate(headers):
    table.cell(0, i).text = text
    shade_cell(table.cell(0, i), PALE_BLUE)
    table.cell(0, i).paragraphs[0].runs[0].bold = True
data = [
    ("Dominion", "Rule the stack", "Spells and abilities", "Counter, protect, redirect, suppress", "Removal and broad taxes"),
    ("Torment", "Make every choice hurt", "Player decisions and resources", "Pay, reveal, choose, mark, risk", "Random or unavoidable punishment"),
    ("Control", "Reshape the battlefield", "Permanents, zones, and combat", "Strip, tap, disable, divide, interdict", "Counterspells and hard locks"),
]
for row_i, row in enumerate(data, 1):
    for col_i, value in enumerate(row):
        table.cell(row_i, col_i).text = value
set_repeat_table_header(table.rows[0])

doc.add_heading("Fast classification test", level=2)
add_bullets(doc, [
    "If its main purpose is deciding whether, how, or against what a spell or ability resolves, it is Dominion.",
    "If the affected player chooses how they pay to continue, it is Torment.",
    "If it changes what permanents, zones, or combatants can do after they exist, it is Control.",
])

# Dominion
doc.add_heading("2. Dominion - Rule the Stack", level=1)
doc.add_paragraph(
    "Dominion is the Host's authority made immediate. It protects the Host's plan, interrupts coordinated player turns, and wins narrow contests on the stack. "
    "Its drama comes from precision: the table can see the intervention, sequence around it, bait it, or answer it."
)
doc.add_heading("Mechanical territory", level=2)
add_bullets(doc, [
    "Counter or redirect specific spells, activated abilities, triggered abilities, or copies.",
    "Protect a Host spell or ability from a player response.",
    "Temporarily suppress the source of an answered ability.",
    "React to repeated or coordinated actions without creating a blanket stack lock.",
])
doc.add_heading("Boundary", level=2)
doc.add_paragraph("Dominion is not unconditional permanent removal, random punishment, permanent silence, or a broad tax on normal play.")

add_command(doc, "Seal the Combo", "Dominion I", "APPROVED PLAYTEST",
            "Reaction: after an activated or triggered ability is put on the stack.",
            "Use the standard Authority cost and encounter limit defined by the equipped Command.",
            "Counter target activated or triggered ability. If its source is a permanent, players cannot activate abilities of that permanent until end of turn.",
            "Answers a precise stack event and briefly suppresses its source.",
            "Players may bait the Command, protect the ability, or sequence around the end-of-turn window.",
            "One source and one visible end-of-turn duration.",
            "Review repeatable mana abilities, planeswalker abilities, and sources with several independent abilities.")

add_command(doc, "Answer the Answer", "Dominion II", "APPROVED PLAYTEST",
            "Reaction: after a player spell that targets a Host-controlled spell or ability is put on the stack.",
            "Use the standard Authority cost and encounter limit defined by the equipped Command.",
            "Counter target spell that targets a Host-controlled spell or ability.",
            "Protects a Host stack action through a narrow, readable counter-war window.",
            "Players can decline to answer, bait the protection, or answer this Command if rules permit.",
            "No persistent state.",
            "Keep the target restriction exact so it cannot become a generic counterspell.")

add_command(doc, "Impose Order", "Dominion concept", "DESIGN CONCEPT",
            "Reaction: after a player creates a copy of a spell or ability.",
            "Proposed once per encounter.",
            "Counter target copy. Until end of turn, that player cannot copy spells or abilities.",
            "Stops multiplication on the stack without answering the original action.",
            "The original remains; players can create the copy after the Command has been spent or answer the Command on the stack.",
            "One player and an end-of-turn restriction.",
            "Confirm interactions with storm, casualty, replicate, triggered copies, and copies that are not cast.")

add_command(doc, "Protected Decree", "Dominion concept", "DESIGN CONCEPT",
            "Reaction: after a Host-controlled spell becomes the target of a player spell or ability.",
            "Proposed once per encounter.",
            "Target Host-controlled spell cannot be countered this turn. You may choose new legal targets for it.",
            "Turns a contested Host action into a narrow assertion of authority.",
            "Players can respond before this resolves, remove all legal targets, or answer the resulting effect rather than the spell.",
            "One spell; duration lasts only while relevant.",
            "Check whether retargeting plus uncounterability creates unavoidable removal or bypasses intended immunities.")

add_command(doc, "Final Edict", "Dominion capstone concept", "DESIGN CONCEPT",
            "Reaction: after a Host-controlled spell or ability is countered.",
            "Proposed once per encounter and high Authority cost.",
            "Copy the countered spell or ability. You may choose new legal targets for the copy.",
            "Lets a specialized Host recover one decisive stack action without creating a permanent lock.",
            "The copy uses the stack and can be answered normally.",
            "One copied object; no persistent state.",
            "Exclude Commands themselves if recursion is possible; test extra-turn, mass-removal, and variable-cost effects.")

# Torment
doc.add_heading("3. Torment - Make Every Choice Hurt", level=1)
doc.add_paragraph(
    "Torment weaponizes ambition, fear, and incomplete information. It does not simply take resources; it presents a consequential choice about which resource, plan, or certainty the player is willing to surrender."
)
doc.add_heading("Mechanical territory", level=2)
add_bullets(doc, [
    "Life payments, discard or reveal choices, and temporary behavior taxes.",
    "Punisher choices, visible marks, and risk-for-reward decisions.",
    "Reactions to overextension, repeated casting, or greedy resource use.",
    "Short, announced dilemmas that resolve without secret bookkeeping.",
])
doc.add_heading("Boundary", level=2)
doc.add_paragraph("Torment is not random discard, unavoidable life loss without a decision or trigger, player elimination, permanent hand denial, or punishment for ordinary play with no viable alternative.")

add_command(doc, "Cruel Choice", "Torment I", "APPROVED PLAYTEST",
            "Reaction: after a player casts their second spell in a turn.",
            "Use the standard Authority cost and encounter limit defined by the equipped Command.",
            "That spell's controller chooses one: pay 5 life; or counter that spell.",
            "The player chooses which resource preserves their plan.",
            "Sequence the important spell first, hold life, accept the counter, or answer the Command.",
            "Only the second-spell trigger and immediate choice.",
            "Verify multiplayer turn ownership, copied spells, free spells, and whether life payment can eliminate a player.")

add_command(doc, "Forced Exposure", "Torment II", "APPROVED PLAYTEST",
            "Priority: during a Host main phase while the stack is empty.",
            "Use the standard Authority cost and encounter limit defined by the equipped Command.",
            "Target player reveals their hand. Choose a nonland card name revealed this way. Spells with that name cost {3} more to cast through the end of that player's next turn.",
            "Information becomes pressure, but the player retains alternative lines.",
            "The player can cast other cards, pay the tax, discard the named card for value, or wait out the duration.",
            "One public card name and a duration ending on a specific player's turn.",
            "Check split cards, cards with alternate names, cost reductions, and repeated copies.")

add_command(doc, "Price of Greed", "Torment concept", "DESIGN CONCEPT",
            "Reaction: after a player draws their second card beyond the first card drawn in their draw step this turn.",
            "Proposed once per encounter.",
            "That player chooses one: discard a card; or lose 4 life.",
            "Pressures extra-card engines through a meaningful, immediate resource choice.",
            "The player can stop drawing before the trigger, keep discard fodder, preserve life, or answer the Command.",
            "Immediate choice; no persistent marker.",
            "Define the draw count precisely and test wheel effects, replacement effects, and draws during another player's turn.")

add_command(doc, "Blood Toll", "Torment concept", "DESIGN CONCEPT",
            "Reaction: after a player casts a spell from a graveyard or exile.",
            "Proposed once per encounter.",
            "That spell's controller chooses one: pay life equal to its mana value; or return it to its previous zone.",
            "Makes recursion and impulse casting costly without simply forbidding them.",
            "Choose a lower-cost spell, preserve life, accept the return, or answer the Command.",
            "Immediate choice; remember the previous zone only during resolution.",
            "Define X, lands played from exile, commanders, flashback exile replacement, and zero-mana spells.")

add_command(doc, "Burden of Knowledge", "Torment concept", "DESIGN CONCEPT",
            "Priority: during a Host main phase while the stack is empty.",
            "Proposed once per encounter.",
            "Target player reveals their hand, then chooses a nonland card in it. Exile that card until the beginning of their next turn, or that player loses life equal to twice its mana value.",
            "The player decides whether time or life is the more valuable resource.",
            "Reveal a low-impact card, accept a temporary delay, or pay life to keep the line available.",
            "One face-up exiled card with a clear return trigger.",
            "Prevent choosing an impossible card state; test commanders, adventure cards, and cards with no mana cost.")

add_command(doc, "Marked for Death", "Torment capstone concept", "DESIGN CONCEPT",
            "Reaction: after a creature is declared as an attacker against the Host.",
            "Proposed once per encounter and high Authority cost.",
            "Mark that creature. At the end of combat, its controller chooses one: sacrifice it; or lose 6 life. Remove the mark.",
            "Turns aggression into a visible, delayed dilemma rather than stopping combat outright.",
            "Remove or blink the creature, prevent it from surviving combat, preserve life, or answer the trigger.",
            "Requires one physical mark until end of combat.",
            "Test tokens, commanders, indestructible, sacrifice prevention, combat-ending effects, and player elimination.")

# Control
doc.add_heading("4. Control - Reshape the Battlefield", level=1)
doc.add_paragraph(
    "Control dictates where action is safe and which protections matter. It closes zones, strips shields, and reshapes combat for a short, readable window. It disrupts a plan without erasing the player's ability to play."
)
doc.add_heading("Mechanical territory", level=2)
add_bullets(doc, [
    "Strip hexproof, indestructible, ward, evasion, or another battlefield protection.",
    "Tap, disable, isolate, or restrict a permanent for a short duration.",
    "Interdict graveyards or another clearly defined zone.",
    "Change combat permissions or separate coordinated permanents.",
])
doc.add_heading("Boundary", level=2)
doc.add_paragraph("Control is not a generic counterspell, unrestricted destroy or exile, permanent theft, land denial, turn denial, or a hard lock.")

add_command(doc, "Strip the Shield", "Control I", "APPROVED PLAYTEST",
            "Priority: during a legal Host priority window with a visible battlefield target.",
            "Use the standard Authority cost and encounter limit defined by the equipped Command.",
            "Target player-controlled nonland permanent loses hexproof and indestructible until end of turn.",
            "Changes how an existing permanent can be interacted with.",
            "Protect it by other means, use it before the window, remove the Host's follow-up, or let the duration expire.",
            "One target and an end-of-turn duration.",
            "Confirm layers, ward, shroud, protection, and effects that regain abilities later in the turn.")

add_command(doc, "Grave Interdict", "Control II", "APPROVED PLAYTEST",
            "Priority: during a legal Host priority window.",
            "Use the standard Authority cost and encounter limit defined by the equipped Command.",
            "Exile target card from a graveyard. Until end of turn, players cannot cast spells from graveyards or activate abilities of cards in graveyards.",
            "Closes a specific zone for a short, announced window.",
            "Use graveyard resources first, answer the Command, or shift to hand and battlefield lines.",
            "One exiled card and one global end-of-turn restriction.",
            "Test triggered abilities in graveyards, replacement effects, escape, flashback, and cards that move themselves as a cost.")

add_command(doc, "Arrest the Vanguard", "Control concept", "DESIGN CONCEPT",
            "Reaction: after attackers are declared and before blockers.",
            "Proposed once per encounter.",
            "Tap target attacking creature and remove it from combat. It does not untap during its controller's next untap step.",
            "Reshapes combat through a visible battlefield restriction.",
            "Attack with multiple threats, protect against tapping, untap it later, or choose a different combat window.",
            "One creature and one missed untap.",
            "Test vigilance, already tapped attackers, extra combats, untap-step skipping, and commanders.")

add_command(doc, "Ground the Assault", "Control concept", "DESIGN CONCEPT",
            "Reaction: after attackers are declared.",
            "Proposed once per encounter.",
            "Attacking creatures lose flying and trample until end of turn.",
            "Changes combat geometry without preventing the attack or destroying attackers.",
            "Sequence noncombat damage, use other evasion, remove blockers, or answer the Command.",
            "One group effect through end of turn.",
            "Check ability-removal layers, creatures that regain abilities, and whether the effect is too weak without blockers.")

add_command(doc, "Quarantine", "Control concept", "DESIGN CONCEPT",
            "Priority: during a Host main phase while the stack is empty.",
            "Proposed once per encounter.",
            "Until end of turn, target nonland permanent cannot attack, block, or activate abilities. Its controller may pay {3} to end this effect.",
            "Temporarily closes one battlefield piece while preserving a visible release valve.",
            "Pay the release cost, use mana abilities first, remove or blink the permanent, or wait out the turn.",
            "One target and an end-of-turn duration.",
            "Clarify whether the payment is a special action, how often it may be attempted, and interactions with mana abilities.")

add_command(doc, "Divide and Conquer", "Control capstone concept", "DESIGN CONCEPT",
            "Priority: at the beginning of combat on the players' turn.",
            "Proposed once per encounter and high Authority cost.",
            "Choose up to one creature controlled by each player. Until end of turn, those creatures cannot attack or block together: for each combat, only one of them may be declared as an attacker or blocker.",
            "Breaks coordinated battlefields by forcing the table to choose which threat participates.",
            "Choose the best participant, create additional combats, remove a chosen creature, or answer the Command.",
            "A small visible group of selected creatures through end of turn.",
            "The restriction needs precise multiplayer rules wording; test team turns, forced attacks, extra combats, and more than two players.")

# Framework
doc.add_heading("5. Progression Framework", level=1)
doc.add_paragraph(
    "The long-term target is a broad catalogue without presenting thirty equal choices at once. Rank establishes competence; branches establish play style; achievements provide unusual trophies and build-around options."
)
table = doc.add_table(rows=4, cols=4)
table.style = "Table Grid"
set_table_fixed(table, [1.35, 0.8, 2.05, 2.3])
for i, text in enumerate(["Source", "Count", "Role", "Recommended presentation"]):
    table.cell(0, i).text = text
    shade_cell(table.cell(0, i), PALE_BLUE)
    table.cell(0, i).paragraphs[0].runs[0].bold = True
for i, row in enumerate([
    ("Rank", "6", "Baseline tools and onboarding", "Linear track; automatically unlocked"),
    ("Branches", "18", "Specialization and play style", "Three tabs; six visible nodes per branch"),
    ("Achievements", "6", "Rare, thematic sidegrades", "Separate tab with source badges"),
], 1):
    for j, value in enumerate(row):
        table.cell(i, j).text = value
set_repeat_table_header(table.rows[0])

doc.add_heading("Recommended six-node branch shape", level=2)
add_bullets(doc, [
    "Root: introduces the branch's defining verb.",
    "Foundation: improves reliability or opens the branch's second verb.",
    "Specialization fork A and B: mutually exclusive expressions of the same philosophy.",
    "Mastery: rewards commitment and combines earlier lessons.",
    "Capstone: a dramatic once-per-encounter expression with normal counterplay.",
])
doc.add_paragraph(
    "A complete route should use five points, leaving room for cross-branch experimentation. Avoid nodes that only reduce Authority cost or add uses; those become automatic best picks and weaken the sidegrade model."
)

doc.add_heading("UI direction for the expanded tree", level=2)
add_bullets(doc, [
    "Global screen-space tree browser with tabs: All, Rank, Dominion, Torment, Control, and Achievements.",
    "Scrollable focused view with persistent connectors, current rank, available talent points, and source badges.",
    "Every node shows locked, available, selected, and equipped states without relying on color alone.",
    "Hover or focus reveals a tooltip; selection opens a persistent detail panel with rules text, timing, prerequisites, and counterplay notes.",
    "Command equipping remains on the physical Host Profile object, six Commands per page, matching the Player Profile's Essence-page interaction model.",
])

doc.add_heading("Equipped Command token bags", level=2)
add_bullets(doc, [
    "Each equipped Command exposes a Spawn control on its loadout entry.",
    "Spawn creates the configured Infinite Bag for that Command; the bag contains its custom Command card/token.",
    "A placeholder bag is acceptable during development, but the Command definition should reserve a stable asset key.",
    "The loadout UI should distinguish unlocked, equipped, and spawned-object state; spawning never changes unlock or equip state.",
])

# Evaluation rubric
doc.add_heading("6. Command Evaluation Rubric", level=1)
doc.add_paragraph("A proposal should pass all seven questions before implementation:")
questions = [
    ("Branch identity", "Would a designer classify it correctly without seeing the branch label?"),
    ("Timing clarity", "Can the Host and players identify the exact legal window without debate?"),
    ("Agency", "Can players sequence, pay, choose, answer, or avoid the effect in a meaningful way?"),
    ("Table load", "Can the effect be tracked with zero or one visible token and a clear duration?"),
    ("Resolution speed", "Does it resolve in one short exchange rather than opening a nested minigame?"),
    ("Repeatability", "Is it exciting once per encounter but safe from loops, recursion, and repeated denial?"),
    ("Rules coverage", "Have multiplayer, copy, commander, zone-change, replacement, and layer interactions been considered?"),
]
table = doc.add_table(rows=len(questions) + 1, cols=2)
table.style = "Table Grid"
set_table_fixed(table, [1.45, 5.05])
table.cell(0, 0).text = "Check"
table.cell(0, 1).text = "Pass condition"
for cell in table.rows[0].cells:
    shade_cell(cell, PALE_BLUE)
    cell.paragraphs[0].runs[0].bold = True
for i, (a, b) in enumerate(questions, 1):
    table.cell(i, 0).text = a
    table.cell(i, 1).text = b
set_repeat_table_header(table.rows[0])

doc.add_heading("Balance levers, in preferred order", level=2)
add_bullets(doc, [
    "Narrow the trigger or legal targets.",
    "Strengthen the player's counterplay or choice.",
    "Shorten the duration.",
    "Limit the Command to once per encounter.",
    "Increase Authority cost only after the behavior is healthy; cost should not excuse an unfun effect.",
])

doc.add_heading("Playtest observations to capture", level=2)
add_bullets(doc, [
    "Was the trigger noticed immediately by everyone?",
    "Did the player understand their options before choosing?",
    "Did the Command create a new line or merely delete the current one?",
    "How long did resolution take?",
    "Did the table forget any duration or marker?",
    "Would the effect remain fun if seen in three consecutive games?",
])

# Template
doc.add_heading("7. Command Proposal Worksheet", level=1)
doc.add_paragraph("Copy this section for each new proposal. Write rules text only after the branch promise, timing, and counterplay are clear.")
fields = [
    ("Name", "Short, evocative, and distinct from existing cards or system terms."),
    ("Status", "Design concept / rules review / approved playtest / implemented."),
    ("Branch and node", "Dominion, Torment, or Control; root, foundation, fork, mastery, or capstone."),
    ("Fantasy sentence", "One sentence describing what the Host feels like they are doing."),
    ("Timing", "Exact trigger or priority window."),
    ("Cost and limit", "Authority cost, once-per-turn/encounter restriction, and target restrictions."),
    ("Rules text", "Table-ready behavior using established MTG language where possible."),
    ("Branch identity", "Why this belongs here and does not belong in the other two branches."),
    ("Counterplay", "At least one realistic sequencing, payment, choice, or response line."),
    ("Tracking", "Tokens, named cards, duration, and who is responsible for remembering it."),
    ("Abuse check", "Copies, recursion, extra turns/combats, commanders, zone changes, and multiplayer scaling."),
    ("Asset key", "Stable identifier for the eventual Command card and Infinite Bag configuration."),
]
table = doc.add_table(rows=len(fields) + 1, cols=2)
table.style = "Table Grid"
set_table_fixed(table, [1.55, 4.95])
table.cell(0, 0).text = "Field"
table.cell(0, 1).text = "Required content"
for cell in table.rows[0].cells:
    shade_cell(cell, PALE_BLUE)
    cell.paragraphs[0].runs[0].bold = True
for i, (a, b) in enumerate(fields, 1):
    table.cell(i, 0).text = a
    table.cell(i, 1).text = b
set_repeat_table_header(table.rows[0])

doc.add_heading("Definition of ready for implementation", level=2)
add_bullets(doc, [
    "Rules text and timing are approved by rules design.",
    "Authority cost and encounter limit are assigned.",
    "Unlock source, prerequisite, and equip eligibility are defined.",
    "UI title, short description, full tooltip, and asset key are supplied.",
    "Token/bag behavior is defined, even if placeholder art is used.",
    "At least one automated or scripted test covers legality and one covers a denied/invalid use.",
])

# References
doc.add_heading("8. Design References and Ownership", level=1)
doc.add_paragraph(
    "The progression presentation may borrow high-level lessons from game systems such as Total War: Warhammer III lord/technology trees, "
    "Final Fantasy job or license-style specialization, and Skyrim's perk constellations: visible paths, thematic clusters, prerequisites, and clear node details. "
    "MTGR's implementation should remain purpose-built for a shared tabletop session and its lower tolerance for hidden state."
)
doc.add_heading("Canonical project references", level=2)
refs = [
    ("Host Skill Tree rules", r"D:\Github\MTGR-Rules\rules\host\host-skill-tree.md"),
    ("Host profile rules", r"D:\Github\MTGR-Rules\rules\host\host-profile.md"),
    ("Host expectations", r"D:\Github\MTGR-Rules\rules\host\host-expectations.md"),
]
for label, path in refs:
    p = doc.add_paragraph()
    p.add_run(label + ": ").bold = True
    p.add_run(path)

add_status_box(
    doc,
    "Ownership note",
    "The approved rules files remain canonical. This guide combines those decisions with unapproved examples for development discussion; concept text must not be treated as live game rules until promoted through review.",
    PALE_RED,
)

doc.core_properties.title = "MTGR Host Command Branch Design Guide"
doc.core_properties.subject = "Branch philosophy, progression framework, and example Host Commands"
doc.core_properties.author = "MTGR Development Team"
doc.core_properties.keywords = "MTGR, Host, Commands, Dominion, Torment, Control, progression"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
