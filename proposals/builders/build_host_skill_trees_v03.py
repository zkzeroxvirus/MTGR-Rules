from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "MTGR-HOST-SKILL-TREES-REWORK-V0.3.docx"

NAVY = "1F355E"
BLUE = "2E74B5"
PURPLE = "7030A0"
RED = "9C0006"
GREEN = "375623"
GOLD = "806000"
INK = "202124"
MUTED = "666666"
LIGHT = "F3F5F8"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=110, bottom=70, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field_page(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_callout(doc, label, text, fill="E8EEF5"):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ": ")
    set_font(r, bold=True, color=NAVY)
    r = p.add_run(text)
    set_font(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_labeled(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(label + ": ")
    set_font(r, bold=True, color=NAVY)
    r = p.add_run(text)
    set_font(r)


def add_tree_path(doc, left, right, mastery, capstone, color):
    table = doc.add_table(rows=4, cols=3)
    set_table_geometry(table, [1750, 5860, 1750])
    rows = [
        ("TIER I", left[0], right[0]),
        ("TIER II", left[1], right[1]),
        ("TIER III", mastery, "Paths converge"),
        ("CAPSTONE", capstone, "Branch culmination"),
    ]
    for i, values in enumerate(rows):
        for j, value in enumerate(values):
            cell = table.cell(i, j)
            cell.text = value
            shade(cell, LIGHT if i < 2 else "E8EEF5")
            for run in cell.paragraphs[0].runs:
                set_font(run, size=9.5, bold=(j == 0 or i >= 2), color=color if j == 0 else INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_command_register(doc, commands, color):
    headers = ["Node / Command", "Unlock & cost", "Timing", "Rules text"]
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [1700, 1750, 1250, 4660])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        shade(cell, color)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9.5, bold=True, color=WHITE)
    set_repeat_header(table.rows[0])
    for idx, command in enumerate(commands):
        tier, name, unlock, cost, timing, effect = command
        cells = table.add_row().cells
        values = [f"{tier}\n{name}", f"{unlock}\n{cost}", timing, effect]
        for j, value in enumerate(values):
            cells[j].text = value
            if idx % 2:
                shade(cells[j], "F8F9FB")
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_font(run, size=8.7, bold=(j == 0), color=color if j == 0 else INK)
    return table


BRANCHES = [
    {
        "name": "Dominion",
        "color": BLUE,
        "philosophy": "Rule the moment. Dominion exploits timing, targets, ownership, and resolution. It normally converts or delays a player action; unconditional denial is reserved for the capstone.",
        "identity": "A Dominion Command should make a pivotal player action resolve differently, later, or for the Host's benefit. Tier I is narrow and opportunistic; deeper nodes become more reliable and comprehensive.",
        "paths": (("Claim the Initiative", "Imperial Reflection"), ("Turn the Blade", "Delay the Inevitable"), "Arcane Supremacy", "Final Authority"),
        "commands": [
            ("Tier I-A", "Claim the Initiative", "Rank 2; 1 TP", "1 Authority; 1/encounter", "Reaction", "Copy target instant or sorcery spell controlled by a player with mana value 3 or less. The Host may choose new targets for the copy."),
            ("Tier I-B", "Turn the Blade", "Rank 2; 1 TP", "1 Authority; 1/encounter", "Reaction", "Change the target of target spell or ability with exactly one target. A different legal target must exist."),
            ("Tier II-A", "Imperial Reflection", "Rank 4; Claim the Initiative; 1 TP", "1 Authority; 1/encounter", "Reaction", "Copy target activated or triggered ability controlled by a player. The Host may choose new targets. If the copy cannot legally benefit the Host or harm a player, the Host instead draws two cards."),
            ("Tier II-B", "Delay the Inevitable", "Rank 4; Turn the Blade; 1 TP", "1 Authority; 1/encounter", "Reaction", "Exile target spell controlled by a player with two time counters. Remove one at the beginning of its owner's upkeep. When the last is removed, its owner may cast it without paying its mana cost; it cannot target the Host or a Host-controlled object that turn."),
            ("Tier III", "Arcane Supremacy", "Rank 7; either Tier II Dominion node; 1 TP", "2 Authority; 1/encounter", "Reaction", "Copy target spell controlled by a player, then return the original to its owner's hand. The Host may choose new targets for the copy. This Command cannot be copied or redirected."),
            ("Capstone", "Final Authority", "Rank 10; Arcane Supremacy; 1 TP", "2 Authority; 1/encounter", "Reaction", "Counter target spell or ability and exile it. Final Authority cannot be countered, copied, or redirected."),
        ],
    },
    {
        "name": "Torment",
        "color": RED,
        "philosophy": "Make the party choose the disaster. Torment presents a bad outcome and a potentially worse alternative, attacking life, cards, permanents, tempo, and cooperation.",
        "identity": "The party chooses collectively; the Host controls whether suffering occurs. Neither option may be a trivial payment. The correct answer should change with life totals, hands, graveyards, and battlefield state.",
        "paths": (("Blood or Tribute", "Kneel or Bleed"), ("Feed the Darkness", "Divide the Alliance"), "No Safe Answer", "The Final Bargain"),
        "commands": [
            ("Tier I-A", "Blood or Tribute", "Rank 2; 1 TP", "1 Authority; 1/encounter", "Host main phase", "The party chooses - Blood: each player loses 3 life with 1-2 players, 4 life with 3-4 players, or 5 life with 5-6 players. Tribute: the Host draws two cards, creates Treasure equal to Party Scale, and creates a 3/3 black Demon token with flying."),
            ("Tier I-B", "Feed the Darkness", "Rank 2; 1 TP", "1 Authority; 1/encounter", "Host main phase", "The party chooses - Starve: each player discards a card. Feed: the Host returns a permanent card from the Host discard pile to hand and creates Treasure equal to Party Scale."),
            ("Tier II-A", "Kneel or Bleed", "Rank 4; Blood or Tribute; 1 TP", "1 Authority; 1/encounter", "Host main phase", "The party chooses - Kneel: tap one creature controlled by each player and put a stun counter on each. Bleed: until the Host's next turn, whenever a player attacks the Host or activates a nonmana ability of a nonland permanent, that player loses 1 life; no more than three triggers per player."),
            ("Tier II-B", "Divide the Alliance", "Rank 4; Feed the Darkness; 1 TP", "1 Authority; 1/encounter", "Host main phase", "The party chooses - Together: each player loses 3 life and discards a card. Alone: the party chooses one player; that player loses 7 life, discards three cards, and sacrifices their highest-mana-value nonland permanent. The Host breaks a tied victim vote."),
            ("Tier III", "No Safe Answer", "Rank 7; either Tier II Torment node; 1 TP", "2 Authority; 1/encounter", "Host main phase", "The party chooses - Ruin: each player sacrifices their highest-mana-value nonland permanent and the Host draws a card for each player unable to do so. Despair: each player discards a card, mills three cards, and loses 4 life."),
            ("Capstone", "The Final Bargain", "Rank 10; No Safe Answer; 1 TP", "2 Authority; 1/encounter", "Host priority", "The party chooses - Submit: end the current turn and the Host takes an additional turn. Resist: each player loses 5 life and sacrifices their two highest-mana-value nonland permanents. The Final Bargain cannot be countered."),
        ],
    },
    {
        "name": "Suppression",
        "color": GREEN,
        "philosophy": "Break the position already built. Suppression attacks combat formations, engines, graveyards, protections, and established permanents after they reach play.",
        "identity": "Tier I disables a limited number of relevant objects. Deeper specialization widens the affected battlefield and duration until the capstone can suppress the entire party for a turn cycle.",
        "paths": (("Break the Formation", "Seize the Champion"), ("Silence the Engine", "Seal the Threshold"), "Total Lockdown", "Absolute Suppression"),
        "commands": [
            ("Tier I-A", "Break the Formation", "Rank 2; 1 TP", "1 Authority; 1/encounter", "Before combat or after attackers", "Tap creatures controlled by players equal to Party Scale and put a stun counter on each. Host creatures gain menace until end of turn."),
            ("Tier I-B", "Silence the Engine", "Rank 2; 1 TP", "1 Authority; 1/encounter", "Host priority", "Choose nonland permanents controlled by players equal to Party Scale. Until the Host's next turn, they lose all nonmana abilities. If one leaves play during that time, the Host draws a card; draw no more than two cards this way."),
            ("Tier II-A", "Seize the Champion", "Rank 4; Break the Formation; 1 TP", "1 Authority; 1/encounter", "Host main phase", "Gain control of the highest-power creature controlled by target player until the Host's next turn. Untap it. It gains haste and must attack its owner if able."),
            ("Tier II-B", "Seal the Threshold", "Rank 4; Silence the Engine; 1 TP", "1 Authority; 1/encounter", "Host priority", "Exile up to Party Scale cards from player graveyards. Until the Host's next turn, player cards cannot leave graveyards, and creatures entering or dying do not cause player-controlled abilities to trigger."),
            ("Tier III", "Total Lockdown", "Rank 7; either Tier II Suppression node; 1 TP", "2 Authority; 1/encounter", "Host main phase", "Tap all nonland permanents controlled by players. Choose one nonland permanent controlled by each player; those permanents do not untap during their controllers' next untap steps. Until the Host's next turn, players cannot activate nonmana abilities of tapped permanents."),
            ("Capstone", "Absolute Suppression", "Rank 10; Total Lockdown; 1 TP", "2 Authority; 1/encounter", "Host priority", "Until the Host's next turn, player permanents lose hexproof and indestructible, their nonmana abilities cannot be activated, and player-controlled abilities do not trigger. Destroy one nonland permanent controlled by each player. This Command cannot be countered."),
        ],
    },
    {
        "name": "Ascendancy",
        "color": PURPLE,
        "philosophy": "Become the threat. Ascendancy develops the Host's battlefield, converts expendable resources, recovers defeated threats, and turns an established Host position into overwhelming pressure.",
        "identity": "Ascendancy is the proactive alternative to reactive denial. Tier I establishes material; deeper nodes magnify a champion, recur threats, and culminate in a board state the party must answer immediately.",
        "paths": (("Muster the Court", "Crowned Champion"), ("Feed the Throne", "Relentless Return"), "Overwhelming Presence", "The Host Ascendant"),
        "commands": [
            ("Tier I-A", "Muster the Court", "Rank 2; 1 TP", "1 Authority; 1/encounter", "Host main phase", "Create 3/3 Host creature tokens equal to Party Scale. Choose vigilance, menace, or reach; those tokens gain the chosen ability until the encounter ends."),
            ("Tier I-B", "Feed the Throne", "Rank 2; 1 TP", "1 Authority; 1/encounter", "Host main phase", "Draw two cards, then discard a card. The Host may sacrifice a token or Treasure. If they do, create Treasure equal to Party Scale and draw an additional card."),
            ("Tier II-A", "Crowned Champion", "Rank 4; Muster the Court; 1 TP", "1 Authority; 1/encounter", "Beginning of Host combat", "Choose a Host creature. Put two +1/+1 counters on it, untap it, and it gains trample, haste, and ward 2 until the Host's next turn. When it first deals combat damage to a player this turn, draw two cards."),
            ("Tier II-B", "Relentless Return", "Rank 4; Feed the Throne; 1 TP", "1 Authority; 1/encounter", "Host main phase", "Return target Host creature card from the Host discard pile to play with a finality marker. It enters with two +1/+1 counters and gains haste until end of turn."),
            ("Tier III", "Overwhelming Presence", "Rank 7; either Tier II Ascendancy node; 1 TP", "2 Authority; 1/encounter", "Beginning of Host combat", "Until the Host's next turn, Host creatures get +2/+2 and gain trample and ward 2. Whenever one or more Host creatures deal combat damage to a player this turn, create a Treasure and draw a card; trigger once per player."),
            ("Capstone", "The Host Ascendant", "Rank 10; Overwhelming Presence; 1 TP", "2 Authority; 1/encounter", "Host main phase", "Return up to one Host creature from the discard pile to play with a finality marker; create 4/4 Avatar tokens equal to Party Scale; and draw until the Host has five cards in hand. Host creatures gain haste until end of turn. This Command cannot be countered."),
        ],
    },
]


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, color, before, after in (
        ("Title", 27, NAVY, 0, 8),
        ("Subtitle", 13, MUTED, 0, 18),
        ("Heading 1", 18, NAVY, 18, 9),
        ("Heading 2", 14, BLUE, 14, 7),
        ("Heading 3", 11.5, NAVY, 10, 5),
    ):
        st = styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = style_name != "Subtitle"
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "MTGR  |  HOST PROGRESSION DESIGN"
    for run in header.runs:
        set_font(run, size=8.5, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("PLAYTEST PROPOSAL V0.3  |  ")
    set_font(r, size=8.5, color=MUTED)
    add_field_page(footer)

    p = doc.add_paragraph(style="Title")
    p.add_run("Host Skill Trees Rework")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Four specialized Command branches calibrated for 20-life, 40-card players and 1vMany encounters")
    add_callout(doc, "Design objective", "Every equipped Command must materially alter the next turn cycle. Tier I teaches a branch through a narrower effect, while deeper investment increases reliability, scope, and severity.", "E8EEF5")

    doc.add_heading("1. Format and balance assumptions", level=1)
    add_labeled(doc, "Format", "One Host pilots encounter decks against a party. The primary target is 1v6; 1v4 is the next most common configuration. Lower player counts must remain supported.")
    add_labeled(doc, "Player baseline", "Each party member uses a 40-card deck and begins with 20 life. Party-wide discard, mill, life loss, and removal therefore scale quickly and must be evaluated per player as well as in aggregate.")
    add_labeled(doc, "Interaction benchmark", "The legacy benchmark is two free, broad counterspell interactions per encounter. The new system should replace that raw denial with two specialized Commands whose individual effects are worth selecting.")
    add_labeled(doc, "Progression curve", "Tier I is simpler and narrower, not negligible. Tier II changes tactics. Tier III reshapes an encounter. A capstone may be potentially decisive and consumes the Host's normal encounter budget.")

    doc.add_heading("2. Shared scaling and resolution rules", level=1)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2300, 2300, 4760])
    for i, h in enumerate(("Non-Host players", "Party Scale", "Blood loss per player")):
        table.rows[0].cells[i].text = h
        shade(table.rows[0].cells[i], NAVY)
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            set_font(run, bold=True, color=WHITE)
    for values in (("1-2", "1", "3 life"), ("3-4", "2", "4 life"), ("5-6", "3", "5 life")):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
            for run in cells[i].paragraphs[0].runs:
                set_font(run)
    doc.add_paragraph()
    add_labeled(doc, "Party Scale", "Half the number of non-Host players, rounded up. It is used for Host token and Treasure production or for a limited number of player targets. Party-wide effects normally affect each player once and do not multiply again by Party Scale.")
    add_labeled(doc, "Command budget", "Standard nodes cost 1 Authority. Tier III and capstone nodes cost 2 Authority. Every Command is limited to once per encounter. This document assumes at least two potential Command invocations per encounter.")
    add_labeled(doc, "Torment vote", "Players discuss, then vote simultaneously. The option with the most votes occurs; the Host breaks ties. A choice requiring a payment cannot be selected unless the party can complete its mandatory actions.")
    add_labeled(doc, "Finality marker", "A permanent with a finality marker is exiled instead of returning to the Host discard pile when it next leaves play. This limits recursion loops.")

    doc.add_heading("3. Four-branch overview", level=1)
    overview = doc.add_table(rows=1, cols=3)
    set_table_geometry(overview, [1900, 2500, 4960])
    for i, h in enumerate(("Branch", "Primary game space", "Progression identity")):
        overview.rows[0].cells[i].text = h
        shade(overview.rows[0].cells[i], NAVY)
        for run in overview.rows[0].cells[i].paragraphs[0].runs:
            set_font(run, bold=True, color=WHITE)
    summaries = (
        ("Dominion", "Stack, targets, timing", "Conversion and delay progress toward rare absolute denial."),
        ("Torment", "Party decisions and resources", "Simple bad-or-worse choices progress toward encounter-defining collective disasters."),
        ("Suppression", "Battlefield, engines, graveyards", "Limited disablement progresses toward party-wide lockdown."),
        ("Ascendancy", "Host board, recursion, combat", "Material development progresses toward overwhelming Host presence."),
    )
    for row in summaries:
        cells = overview.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for run in cells[i].paragraphs[0].runs:
                set_font(run, bold=(i == 0))

    for idx, branch in enumerate(BRANCHES, start=1):
        doc.add_page_break()
        doc.add_heading(f"4.{idx} {branch['name']} Skill Tree", level=1)
        add_callout(doc, "Philosophy", branch["philosophy"], "F2F4F7")
        add_labeled(doc, "Identity test", branch["identity"])
        left, right, mastery, capstone = branch["paths"]
        add_tree_path(doc, left, right, mastery, capstone, branch["color"])
        doc.add_heading(f"{branch['name']} Command register", level=2)
        add_command_register(doc, branch["commands"], branch["color"])

    doc.add_page_break()
    doc.add_heading("5. Playtest acceptance criteria", level=1)
    criteria = [
        ("Tier I relevance", "A Host can explain why either Tier I node was worth one of only two encounter interactions."),
        ("Progression visibility", "Testers can identify a clear increase in scope, reliability, or severity from Tier I through the capstone."),
        ("Branch identity", "A tester can identify the branch from its effect without seeing the branch label."),
        ("1v4 and 1v6 balance", "No effect becomes negligible at four players or automatically decisive at six solely because it scales once per player."),
        ("Counterspell comparison", "Non-Dominion Hosts voluntarily equip proactive Commands instead of asking to restore two universal free counters."),
        ("Resolution burden", "The Command can be resolved publicly without private notes, extended arithmetic, or unclear expiration timing."),
        ("Fun and counterplay", "Players experience a serious threat while retaining meaningful sequencing, voting, protection, or resource decisions."),
    ]
    for label, text in criteria:
        add_labeled(doc, label, text)
    add_callout(doc, "Playtest status", "All names, values, unlock gates, and effects are provisional. Test Tier I at both 1v4 and 1v6 before implementing every advanced node. Record selection rate, outcome chosen, immediate resource swing, resolution time, and whether the Command changed the next turn cycle.", "FFF2CC")

    doc.core_properties.title = "MTGR Host Skill Trees Rework V0.3"
    doc.core_properties.subject = "Four-branch Host Command specialization proposal"
    doc.core_properties.author = "MTGR Design Team"
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
